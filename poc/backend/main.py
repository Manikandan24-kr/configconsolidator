"""
Kriyadocs Config Consolidator — FastAPI Backend
Sprint 1: New session/rules/export API + preserved QC + Knowledge Agent routes.
"""

import json
import re
import uuid
import os
from dataclasses import asdict
from pathlib import Path
from typing import Optional, List

from dotenv import load_dotenv
from fastapi import FastAPI, File, UploadFile, HTTPException, BackgroundTasks, Depends
from fastapi.middleware.cors import CORSMiddleware
from fastapi.responses import StreamingResponse
from pydantic import BaseModel
from sqlalchemy.orm import Session as DBSession

load_dotenv()

# ── DB + Models ────────────────────────────────────────────────────────────────
from database import create_tables, get_db, check_connection
from models import (
    Customer, Session as SessionModel, UploadedFile, Rule, AuditLog,
    VALID_CATEGORIES, VALID_RULE_STATUSES, _uuid, _now,
)

# ── Extraction pipeline ────────────────────────────────────────────────────────
from extractor.pipeline import run_extraction_pipeline

# ── QC + Knowledge Agent (kept from POC) ──────────────────────────────────────
from quality_checker import run_quality_check, qc_report_to_dict
from sample_manuscript import SAMPLE_MANUSCRIPT

HAS_API_KEY = bool(os.getenv("OPENAI_API_KEY"))

try:
    from rule_extractor import (
        extract_rules_from_section as ai_extract_rules,
        detect_discrepancies,
        rules_to_dict,
        discrepancies_to_dict,
    )
    if not HAS_API_KEY:
        from llm_cache import is_cache_warm
        if is_cache_warm():
            HAS_API_KEY = True
            print("CACHE mode: Warm LLM cache detected — using cached AI results.")
        else:
            raise ImportError("No API key and no warm cache")
    else:
        print("AI mode: OpenAI API key detected.")
except Exception:
    from local_extractor import (
        extract_rules_locally as ai_extract_rules,
        detect_discrepancies_locally as detect_discrepancies,
        rules_to_dict,
        discrepancies_to_dict,
    )
    print("LOCAL mode: Using heuristic extraction.")

UPLOAD_DIR = Path("/tmp/kriyadocs_uploads")
UPLOAD_DIR.mkdir(parents=True, exist_ok=True)

# Legacy in-memory stores (kept for QC/Knowledge Agent tab compatibility)
jobs: dict = {}
extracted_rules_store: dict = {}
discrepancies_store: dict = {}
documents_store: dict = {}

# ── App ────────────────────────────────────────────────────────────────────────
app = FastAPI(title="Kriyadocs Config Consolidator API", version="2.0.0")
app.add_middleware(
    CORSMiddleware,
    allow_origins=["*"],
    allow_methods=["*"],
    allow_headers=["*"],
)

create_tables()


# ── Pydantic schemas ───────────────────────────────────────────────────────────

class CreateSessionRequest(BaseModel):
    customer_name: str
    session_name: str = ""

class UpdateRuleRequest(BaseModel):
    status: Optional[str] = None
    rule: Optional[str] = None
    remarks: Optional[str] = None
    category: Optional[str] = None

class AddRuleRequest(BaseModel):
    category: str
    rule: str
    remarks: Optional[str] = None

class BulkStatusRequest(BaseModel):
    rule_ids: List[str]
    action: str  # "confirm" or "exclude"

class QCRequest(BaseModel):
    text: str
    document_name: str = "Untitled"
    use_ai: bool = True

class KnowledgeQuery(BaseModel):
    question: str
    session_id: str = ""


# ── Health ─────────────────────────────────────────────────────────────────────

@app.get("/api/health")
def health():
    from llm_cache import cache_stats
    stats = cache_stats()
    return {
        "status": "ok",
        "ai_available": HAS_API_KEY,
        "mode": "openai" if HAS_API_KEY else "local",
        "db": "connected" if check_connection() else "error",
        "cache_entries": stats.get("total_entries", 0),
    }


# ── Sessions ───────────────────────────────────────────────────────────────────

@app.get("/api/sessions")
def list_sessions(db: DBSession = Depends(get_db)):
    sessions = db.query(SessionModel).order_by(SessionModel.created_at.desc()).all()
    return [_session_summary(s, db) for s in sessions]


@app.post("/api/sessions", status_code=201)
def create_session(req: CreateSessionRequest, db: DBSession = Depends(get_db)):
    customer = db.query(Customer).filter(Customer.name == req.customer_name).first()
    if not customer:
        customer = Customer(id=_uuid(), name=req.customer_name)
        db.add(customer)
        db.flush()

    name = req.session_name or f"{req.customer_name} — {_now().strftime('%B %Y')}"
    session = SessionModel(
        id=_uuid(),
        customer_id=customer.id,
        name=name,
        status="pending",
    )
    db.add(session)
    db.commit()
    return _session_summary(session, db)


@app.get("/api/sessions/{session_id}")
def get_session(session_id: str, db: DBSession = Depends(get_db)):
    session = _get_session_or_404(session_id, db)
    summary = _session_summary(session, db)

    # Category breakdown
    categories = {}
    for cat in VALID_CATEGORIES:
        rules = db.query(Rule).filter(
            Rule.session_id == session_id,
            Rule.category == cat,
        ).all()
        categories[cat] = {
            "total": len(rules),
            "confirmed": sum(1 for r in rules if r.status == "confirmed"),
            "excluded": sum(1 for r in rules if r.status == "excluded"),
            "pending": sum(1 for r in rules if r.status == "pending"),
        }
    summary["categories"] = categories
    return summary


@app.delete("/api/sessions/{session_id}", status_code=204)
def delete_session(session_id: str, db: DBSession = Depends(get_db)):
    session = _get_session_or_404(session_id, db)
    db.delete(session)
    db.commit()


# ── Upload & Extraction ────────────────────────────────────────────────────────

@app.post("/api/sessions/{session_id}/upload")
async def upload_files(
    session_id: str,
    files: List[UploadFile] = File(...),
    db: DBSession = Depends(get_db),
):
    session = _get_session_or_404(session_id, db)
    session_dir = UPLOAD_DIR / session_id
    session_dir.mkdir(parents=True, exist_ok=True)

    saved = []
    for f in files:
        if not f.filename.lower().endswith(".pdf"):
            continue
        dest = session_dir / f.filename
        content = await f.read()
        dest.write_bytes(content)

        uf = UploadedFile(
            id=_uuid(),
            session_id=session_id,
            filename=f.filename,
            filepath=str(dest),
        )
        db.add(uf)
        saved.append(f.filename)

    session.total_pdfs = len(saved)
    session.status = "uploading"
    db.commit()

    return {"files_uploaded": len(saved), "filenames": saved}


@app.post("/api/sessions/{session_id}/extract")
def start_extraction(
    session_id: str,
    background_tasks: BackgroundTasks,
    db: DBSession = Depends(get_db),
):
    session = _get_session_or_404(session_id, db)
    files = db.query(UploadedFile).filter(UploadedFile.session_id == session_id).all()
    if not files:
        raise HTTPException(400, "No files uploaded for this session")

    session.status = "extracting"
    session.extraction_progress = 0
    session.extraction_message = "Queued for extraction..."
    db.commit()

    background_tasks.add_task(_run_extraction_bg, session_id)
    return {"status": "extracting", "session_id": session_id}


def _run_extraction_bg(session_id: str):
    from database import SessionLocal
    db = SessionLocal()
    try:
        run_extraction_pipeline(session_id, db)
    except Exception as e:
        session = db.query(SessionModel).filter(SessionModel.id == session_id).first()
        if session:
            session.status = "review"
            session.extraction_message = f"Extraction error: {str(e)}"
            db.commit()
    finally:
        db.close()


@app.get("/api/sessions/{session_id}/status")
def get_extraction_status(session_id: str, db: DBSession = Depends(get_db)):
    session = _get_session_or_404(session_id, db)
    files = db.query(UploadedFile).filter(UploadedFile.session_id == session_id).all()
    return {
        "status": session.status,
        "progress": session.extraction_progress,
        "message": session.extraction_message,
        "files_total": len(files),
        "files_done": sum(1 for f in files if f.status == "done"),
        "files_failed": sum(1 for f in files if f.status == "failed"),
        "rules_found_so_far": session.total_rules,
    }


# ── Rules ──────────────────────────────────────────────────────────────────────

@app.get("/api/sessions/{session_id}/rules")
def list_rules(
    session_id: str,
    category: Optional[str] = None,
    status: Optional[str] = None,
    db: DBSession = Depends(get_db),
):
    _get_session_or_404(session_id, db)
    q = db.query(Rule).filter(Rule.session_id == session_id)
    if category:
        q = q.filter(Rule.category == category)
    if status:
        q = q.filter(Rule.status == status)
    rules = q.order_by(Rule.category, Rule.sort_order, Rule.created_at).all()
    return {"rules": [_rule_to_dict(r) for r in rules]}


@app.patch("/api/rules/{rule_id}")
def update_rule(rule_id: str, req: UpdateRuleRequest, db: DBSession = Depends(get_db)):
    rule = db.query(Rule).filter(Rule.id == rule_id).first()
    if not rule:
        raise HTTPException(404, "Rule not found")

    old_val = _rule_to_dict(rule)
    changed = False

    if req.status is not None:
        if req.status not in VALID_RULE_STATUSES:
            raise HTTPException(400, f"Invalid status. Must be one of: {VALID_RULE_STATUSES}")
        rule.status = req.status
        changed = True

    if req.rule is not None:
        rule.rule = req.rule.strip()
        changed = True

    if req.remarks is not None:
        rule.remarks = req.remarks.strip() or None
        changed = True

    if req.category is not None:
        if req.category not in VALID_CATEGORIES:
            raise HTTPException(400, f"Invalid category")
        rule.category = req.category
        changed = True

    if changed:
        rule.updated_at = _now()
        _update_session_counts(rule.session_id, db)

        action = req.status if req.status else "edited"
        log = AuditLog(
            id=_uuid(),
            session_id=rule.session_id,
            rule_id=rule_id,
            action=action,
            old_value=json.dumps(old_val),
            new_value=json.dumps(_rule_to_dict(rule)),
        )
        db.add(log)
        db.commit()

    return _rule_to_dict(rule)


@app.post("/api/sessions/{session_id}/rules", status_code=201)
def add_custom_rule(
    session_id: str,
    req: AddRuleRequest,
    db: DBSession = Depends(get_db),
):
    _get_session_or_404(session_id, db)

    if req.category not in VALID_CATEGORIES:
        raise HTTPException(400, f"Invalid category")

    max_order = db.query(Rule).filter(
        Rule.session_id == session_id,
        Rule.category == req.category,
    ).count()

    rule = Rule(
        id=_uuid(),
        session_id=session_id,
        category=req.category,
        rule=req.rule.strip(),
        source_document="Manual",
        source_section="Manual",
        remarks=req.remarks,
        status="pending",
        is_custom=True,
        sort_order=max_order,
    )
    db.add(rule)

    log = AuditLog(
        id=_uuid(),
        session_id=session_id,
        rule_id=rule.id,
        action="added",
        new_value=json.dumps({"rule": req.rule, "category": req.category}),
    )
    db.add(log)
    db.commit()
    return _rule_to_dict(rule)


@app.delete("/api/rules/{rule_id}", status_code=204)
def delete_rule(rule_id: str, db: DBSession = Depends(get_db)):
    rule = db.query(Rule).filter(Rule.id == rule_id).first()
    if not rule:
        raise HTTPException(404, "Rule not found")
    if not rule.is_custom:
        raise HTTPException(400, "Only custom rules can be deleted. Exclude extracted rules instead.")
    session_id = rule.session_id
    db.delete(rule)
    _update_session_counts(session_id, db)
    db.commit()


@app.post("/api/sessions/{session_id}/rules/bulk")
def bulk_update_rules(
    session_id: str,
    req: BulkStatusRequest,
    db: DBSession = Depends(get_db),
):
    _get_session_or_404(session_id, db)

    if req.action not in ("confirm", "exclude"):
        raise HTTPException(400, "action must be 'confirm' or 'exclude'")

    new_status = "confirmed" if req.action == "confirm" else "excluded"

    rules = db.query(Rule).filter(
        Rule.session_id == session_id,
        Rule.id.in_(req.rule_ids),
    ).all()

    for rule in rules:
        old_status = rule.status
        rule.status = new_status
        rule.updated_at = _now()
        log = AuditLog(
            id=_uuid(),
            session_id=session_id,
            rule_id=rule.id,
            action=f"bulk_{new_status}",
            old_value=json.dumps({"status": old_status}),
            new_value=json.dumps({"status": new_status}),
        )
        db.add(log)

    _update_session_counts(session_id, db)
    db.commit()
    return {"updated": len(rules)}


# ── Export ─────────────────────────────────────────────────────────────────────

@app.get("/api/sessions/{session_id}/export/excel")
def export_excel(
    session_id: str,
    include: str = "confirmed",
    db: DBSession = Depends(get_db),
):
    import io
    import openpyxl
    from openpyxl.styles import Font, PatternFill, Alignment, Border, Side

    session = _get_session_or_404(session_id, db)
    summary = _session_summary(session, db)
    customer_name = summary["customer_name"] or session.name

    q = db.query(Rule).filter(Rule.session_id == session_id)
    if include == "confirmed":
        q = q.filter(Rule.status == "confirmed")
    rules = q.order_by(Rule.category, Rule.sort_order).all()

    rules_by_cat: dict[str, List[Rule]] = {}
    for r in rules:
        rules_by_cat.setdefault(r.category, []).append(r)

    wb = openpyxl.Workbook()
    wb.remove(wb.active)

    header_fill = PatternFill("solid", fgColor="1F3864")
    header_font = Font(bold=True, color="FFFFFF", name="Calibri", size=11)
    alt_fill = PatternFill("solid", fgColor="F2F7FC")
    thin = Side(style="thin", color="B8CCE4")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)

    for cat in VALID_CATEGORIES:
        cat_rules = rules_by_cat.get(cat, [])
        if not cat_rules:
            continue

        sheet_name = cat.replace("_", " ").title()[:31]
        ws = wb.create_sheet(title=sheet_name)
        ws.sheet_view.showGridLines = False

        headers = ["#", "Requirement", "Source Document", "Source Section", "Remarks", "Status"]
        col_widths = [5, 80, 30, 30, 30, 14]
        for ci, (h, w) in enumerate(zip(headers, col_widths), 1):
            ws.column_dimensions[openpyxl.utils.get_column_letter(ci)].width = w
            c = ws.cell(row=1, column=ci, value=h)
            c.fill = header_fill
            c.font = header_font
            c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
            c.border = border
        ws.row_dimensions[1].height = 22

        for ri, rule in enumerate(cat_rules, start=2):
            row_fill = PatternFill("solid", fgColor="FFFFFF") if ri % 2 == 0 else alt_fill
            values = [
                ri - 1,
                rule.rule,
                rule.source_document,
                rule.source_section,
                rule.remarks or "",
                "Custom" if rule.is_custom else rule.status.title(),
            ]
            for ci, val in enumerate(values, 1):
                c = ws.cell(row=ri, column=ci, value=val)
                c.fill = row_fill
                c.font = Font(name="Calibri", size=10)
                c.alignment = Alignment(vertical="top", wrap_text=True)
                c.border = border
            ws.row_dimensions[ri].height = 36

    if not wb.sheetnames:
        ws = wb.create_sheet("No Rules")
        ws["A1"] = "No confirmed rules found."

    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)

    date_str = _now().strftime("%Y%m%d")
    safe_name = re.sub(r"[^\w\-]", "_", customer_name)
    filename = f"{safe_name}_ruleset_{date_str}.xlsx"

    return StreamingResponse(
        buf,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f'attachment; filename="{filename}"'},
    )


@app.get("/api/sessions/{session_id}/export/json")
def export_json(
    session_id: str,
    include: str = "confirmed",
    db: DBSession = Depends(get_db),
):
    import io

    session = _get_session_or_404(session_id, db)
    summary = _session_summary(session, db)
    customer_name = summary["customer_name"] or session.name

    q = db.query(Rule).filter(Rule.session_id == session_id)
    if include == "confirmed":
        q = q.filter(Rule.status == "confirmed")
    rules = q.order_by(Rule.category, Rule.sort_order).all()

    rules_by_cat: dict = {}
    for r in rules:
        rules_by_cat.setdefault(r.category, []).append({
            "id": r.id,
            "rule": r.rule,
            "source_document": r.source_document,
            "source_section": r.source_section,
            "remarks": r.remarks,
            "status": r.status,
            "is_custom": r.is_custom,
        })

    payload = {
        "customer": customer_name,
        "generated": _now().isoformat(),
        "total_rules": sum(len(v) for v in rules_by_cat.values()),
        "categories": rules_by_cat,
    }

    date_str = _now().strftime("%Y%m%d")
    safe_name = re.sub(r"[^\w\-]", "_", customer_name)
    filename = f"{safe_name}_ruleset_{date_str}.json"
    content = json.dumps(payload, indent=2)

    return StreamingResponse(
        io.BytesIO(content.encode()),
        media_type="application/json",
        headers={"Content-Disposition": f'attachment; filename="{filename}"'},
    )


# ── Audit Log ──────────────────────────────────────────────────────────────────

@app.get("/api/sessions/{session_id}/audit")
def get_audit_log(session_id: str, db: DBSession = Depends(get_db)):
    _get_session_or_404(session_id, db)
    logs = (
        db.query(AuditLog)
        .filter(AuditLog.session_id == session_id)
        .order_by(AuditLog.created_at.desc())
        .all()
    )
    return [
        {
            "id": l.id,
            "rule_id": l.rule_id,
            "action": l.action,
            "old_value": json.loads(l.old_value) if l.old_value else None,
            "new_value": json.loads(l.new_value) if l.new_value else None,
            "created_at": l.created_at.isoformat() if l.created_at else None,
        }
        for l in logs
    ]


# ── Legacy Consolidator (kept for QC/Knowledge Agent session compatibility) ────

@app.post("/api/consolidator/upload")
async def legacy_upload(files: List[UploadFile] = File(...)):
    session_id = str(uuid.uuid4())[:8]
    session_dir = UPLOAD_DIR / session_id
    session_dir.mkdir(parents=True, exist_ok=True)
    saved = []
    for f in files:
        if f.filename.lower().endswith(".pdf"):
            dest = session_dir / f.filename
            dest.write_bytes(await f.read())
            saved.append(str(dest))
    jobs[session_id] = {"status": "uploaded", "progress": 0, "message": "Ready", "files": saved}
    return {"session_id": session_id, "files_uploaded": len(saved)}


@app.post("/api/consolidator/process/{session_id}")
async def legacy_process(session_id: str, background_tasks: BackgroundTasks):
    if session_id not in jobs:
        raise HTTPException(404, "Session not found")
    background_tasks.add_task(_legacy_process_session, session_id)
    return {"status": "processing", "session_id": session_id}


def _legacy_process_session(session_id: str):
    from pdf_parser import extract_pdf
    job = jobs.get(session_id, {})
    files = job.get("files", [])
    jobs[session_id] = {**job, "status": "processing", "progress": 5, "message": "Starting..."}
    all_rules, all_discs, all_docs = [], [], []
    for i, fp in enumerate(files):
        jobs[session_id]["progress"] = int(10 + (i / max(len(files), 1)) * 80)
        jobs[session_id]["message"] = f"Processing {Path(fp).name}..."
        doc = extract_pdf(fp)
        all_docs.append(doc)
        for section in doc.sections:
            rules = ai_extract_rules(doc.filename, section.title, section.content)
            all_rules.extend(rules)
    discs = detect_discrepancies(all_rules)
    all_discs.extend(discs)
    extracted_rules_store[session_id] = all_rules
    discrepancies_store[session_id] = all_discs
    documents_store[session_id] = all_docs
    jobs[session_id] = {**job, "status": "complete", "progress": 100, "message": "Done"}


@app.get("/api/consolidator/status/{session_id}")
def legacy_status(session_id: str):
    job = jobs.get(session_id)
    if not job:
        raise HTTPException(404, "Session not found")
    return job


@app.get("/api/consolidator/results/{session_id}")
def legacy_results(session_id: str):
    rules = extracted_rules_store.get(session_id)
    if not rules:
        raise HTTPException(404, "Results not found")
    discs = discrepancies_store.get(session_id, [])
    docs = documents_store.get(session_id, [])
    by_category: dict = {}
    for r in rules:
        cat = r.category
        sub = r.subcategory
        by_category.setdefault(cat, {}).setdefault(sub, []).append(asdict(r))
    return {
        "session_id": session_id,
        "stats": {"total_rules": len(rules), "documents_processed": len(docs)},
        "rules_by_category": by_category,
        "discrepancies": discrepancies_to_dict(discs),
    }


# ── QC (kept from POC) ─────────────────────────────────────────────────────────

@app.post("/api/qc/check")
async def run_qc_check(request: QCRequest):
    if not request.text.strip():
        raise HTTPException(400, "No text provided")
    use_ai = request.use_ai and HAS_API_KEY
    report = run_quality_check(request.text, request.document_name, use_ai=use_ai)
    return qc_report_to_dict(report)


@app.get("/api/qc/sample")
def get_sample():
    return {"text": SAMPLE_MANUSCRIPT, "name": "Sample: Multi-tissue CRISPR Gene Editing Study"}


@app.post("/api/qc/check-sample")
def check_sample():
    report = run_quality_check(SAMPLE_MANUSCRIPT, "Sample Manuscript", use_ai=HAS_API_KEY)
    return qc_report_to_dict(report)


# ── Knowledge Agent (kept from POC) ───────────────────────────────────────────

@app.post("/api/knowledge/ask")
async def ask_knowledge(request: KnowledgeQuery):
    rules = extracted_rules_store.get(request.session_id, [])
    if not rules:
        return {"answer": "No style guide processed yet. Use the Consolidator tab first.", "sources": []}

    if not HAS_API_KEY:
        return _local_knowledge_search(request.question, rules)

    from llm_cache import call_llm
    rules_context = "\n".join(
        f"- [{r.rule_id}] {r.rule}" for r in rules
    )[:50000]
    system = "You are a style guide assistant. Answer based ONLY on the rules provided. Cite rule_ids."
    user = f"RULES:\n{rules_context}\n\nQUESTION: {request.question}"
    try:
        answer = call_llm(system, user, max_tokens=1024)
    except RuntimeError:
        return _local_knowledge_search(request.question, rules)

    sources = [
        {"rule_id": r.rule_id, "rule": r.rule, "source": r.source}
        for r in rules if r.rule_id in answer
    ][:5]
    return {"answer": answer, "sources": sources}


def _local_knowledge_search(question: str, rules: list) -> dict:
    stop = {"how","what","when","where","why","the","a","an","in","on","to","of","and","or","is","are","do","does"}
    keywords = [w for w in re.findall(r'\b[a-z]+\b', question.lower()) if w not in stop and len(w) > 2]
    scored = []
    for r in rules:
        score = sum(3 if kw in r.rule.lower() else 0 for kw in keywords)
        if score > 0:
            scored.append((score, r))
    scored.sort(key=lambda x: -x[0])
    if not scored:
        return {"answer": "No matching rules found. Try different keywords.", "sources": []}
    parts = ["Based on the style rules:\n"]
    for _, r in scored[:8]:
        parts.append(f"**[{r.rule_id}]** {r.rule}")
    return {"answer": "\n".join(parts), "sources": []}


# ── Demo endpoints (kept from POC) ────────────────────────────────────────────

@app.post("/api/demo/load-science-guides")
async def load_science_guides(background_tasks: BackgroundTasks):
    project_root = Path(__file__).parent.parent.parent
    pdf_files = sorted(project_root.glob("*.pdf"))
    if not pdf_files:
        raise HTTPException(404, "No PDF files found in project directory")
    session_id = "science-demo"
    session_dir = UPLOAD_DIR / session_id
    session_dir.mkdir(exist_ok=True)
    file_paths = [str(p) for p in pdf_files]
    jobs[session_id] = {"status": "processing", "progress": 0, "message": "Loading Science guides...", "files": file_paths}
    background_tasks.add_task(_legacy_process_session, session_id)
    return {"session_id": session_id, "files_found": len(file_paths)}


@app.post("/api/demo/load-bmj-guides")
async def load_bmj_guides(background_tasks: BackgroundTasks):
    project_root = Path(__file__).parent.parent.parent
    bmj_dir = project_root / "bmj"
    pdf_files = sorted(bmj_dir.glob("*.pdf"))
    if not pdf_files:
        raise HTTPException(404, "No PDF files found in bmj/ directory")
    session_id = "bmj-demo"
    session_dir = UPLOAD_DIR / session_id
    session_dir.mkdir(exist_ok=True)
    file_paths = [str(p) for p in pdf_files]
    jobs[session_id] = {"status": "processing", "progress": 0, "message": "Loading BMJ guides...", "files": file_paths}
    background_tasks.add_task(_legacy_process_session, session_id)
    return {"session_id": session_id, "files_found": len(file_paths)}


# ── Helpers ────────────────────────────────────────────────────────────────────

def _get_session_or_404(session_id: str, db: DBSession) -> SessionModel:
    session = db.query(SessionModel).filter(SessionModel.id == session_id).first()
    if not session:
        raise HTTPException(404, f"Session '{session_id}' not found")
    return session


def _session_summary(session: SessionModel, db: DBSession) -> dict:
    customer = db.query(Customer).filter(Customer.id == session.customer_id).first()
    return {
        "id": session.id,
        "customer_name": customer.name if customer else "",
        "session_name": session.name,
        "status": session.status,
        "total_pdfs": session.total_pdfs,
        "total_rules": session.total_rules,
        "confirmed_rules": session.confirmed_rules,
        "excluded_rules": session.excluded_rules,
        "extraction_progress": session.extraction_progress,
        "extraction_message": session.extraction_message,
        "created_at": session.created_at.isoformat() if session.created_at else None,
        "updated_at": session.updated_at.isoformat() if session.updated_at else None,
    }


def _rule_to_dict(rule: Rule) -> dict:
    return {
        "id": rule.id,
        "session_id": rule.session_id,
        "category": rule.category,
        "rule": rule.rule,
        "source_document": rule.source_document,
        "source_section": rule.source_section,
        "remarks": rule.remarks,
        "status": rule.status,
        "is_custom": rule.is_custom,
        "sort_order": rule.sort_order,
        "created_at": rule.created_at.isoformat() if rule.created_at else None,
        "updated_at": rule.updated_at.isoformat() if rule.updated_at else None,
    }


def _update_session_counts(session_id: str, db: DBSession) -> None:
    rules = db.query(Rule).filter(Rule.session_id == session_id).all()
    session = db.query(SessionModel).filter(SessionModel.id == session_id).first()
    if session:
        session.total_rules = len(rules)
        session.confirmed_rules = sum(1 for r in rules if r.status == "confirmed")
        session.excluded_rules = sum(1 for r in rules if r.status == "excluded")


if __name__ == "__main__":
    import uvicorn
    uvicorn.run(app, host="0.0.0.0", port=8001)
